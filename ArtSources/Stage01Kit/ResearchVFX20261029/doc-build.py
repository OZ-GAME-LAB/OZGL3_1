from pathlib import Path
import json, os, re, sys, importlib.util, subprocess
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT=Path(__file__).resolve().parents[3]
DATA=Path(__file__).resolve().parent
OUT=ROOT/'planning'/'Project_M_무료VFX_조사와적용계획_2026-10-29.docx'
QA=DATA/'DocQA'
RUNTIME=Path(r'C:/Program Files/WindowsApps/OpenAI.CodexPrimaryRuntime.v26-1017-2736-0_26.1017.2736.0_x64__3k8sg7r9htsxt/dependencies')
SKILL=Path(r'C:/Users/me/.codex/plugins/cache/openai-primary-runtime/documents/26.1017.22136/skills/documents')

def load(name):
    return json.loads((DATA/name).read_text(encoding='utf-8-sig'))

def all_entries(data):
    result=[]
    if isinstance(data,dict):
        if isinstance(data.get('name',data.get('title')),str) and isinstance(data.get('url'),str):
            row=dict(data); row['name']=data.get('name',data.get('title')); result.append(row)
        for v in data.values(): result.extend(all_entries(v))
    elif isinstance(data,list):
        for v in data: result.extend(all_entries(v))
    return result

inputs={name:load(name) for name in ['project-audit.json','supplementary-research.json','unity-research.json','external-research.json','strategy.json']}
entries=all_entries(inputs)

def source(needle):
    found=[e for e in entries if needle.lower() in e.get('name','').lower()]
    if not found: raise ValueError('Missing required verified source: '+needle)
    if needle=='Particle Pack': return next(e for e in found if e.get('creator')=='Kenney')
    # For assets, prefer the verified Asset Store acquisition/license path.
    return next((e for e in found if 'assetstore.unity.com' in e.get('url','')),found[0])

doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.5); sec.page_height=Inches(11)
sec.top_margin=Inches(.70); sec.bottom_margin=Inches(.75)
sec.left_margin=Inches(.72); sec.right_margin=Inches(.72)
sec.footer_distance=Inches(.25)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','List Bullet','List Number','Caption']:
    st=doc.styles[name]; st.font.name='맑은 고딕'; st.font.color.rgb=RGBColor(0,0,0)
    st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'맑은 고딕')
    st._element.rPr.rFonts.set(qn('w:ascii'),'Malgun Gothic')
    st._element.rPr.rFonts.set(qn('w:hAnsi'),'Malgun Gothic')
    st.font.size=Pt(11)
    st.paragraph_format.space_after=Pt(6)
    st.paragraph_format.line_spacing=Pt(16)
normal=doc.styles['Normal']; normal.font.size=Pt(11.3)
normal.paragraph_format.widow_control=True
for name,size in [('Title',24),('Heading 1',17),('Heading 2',12.5)]:
    st=doc.styles[name]; st.font.size=Pt(size); st.font.bold=True
    st.paragraph_format.space_before=Pt(10 if name!='Title' else 0)
    st.paragraph_format.space_after=Pt(8)
    st.paragraph_format.keep_with_next=True
doc.styles['Title'].paragraph_format.line_spacing=Pt(30)
doc.styles['Heading 1'].paragraph_format.line_spacing=Pt(23)
doc.styles['Heading 2'].paragraph_format.line_spacing=Pt(18)
doc.styles['Subtitle'].font.size=Pt(11)
# The default Word theme can contain a blue Title paragraph border.
for border in list(doc.styles.element.xpath('.//w:pBdr')):
    border.getparent().remove(border)
doc.core_properties.title='Project M 무료 VFX 조사와 적용 계획'
doc.core_properties.subject='Unity 6 URP 기반 2.5D 치비 액션의 무료 전투 효과 선정'
doc.core_properties.author='Project M'
doc.core_properties.keywords='VFX, Unity, URP, 무료 에셋, 조사, 적용 계획'

footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run(); r.font.size=Pt(9)
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); r._r.append(fld)

def p(text='',bold=False):
    par=doc.add_paragraph(); run=par.add_run(text); run.bold=bold; return par
def h(text,level=1): return doc.add_heading(text,level=level)
def bullet(text): return doc.add_paragraph(text,style='List Bullet')
def page(title):
    doc.add_page_break(); h(title)
def link(par,label,url):
    rid=par.part.relate_to(url,RT.HYPERLINK,is_external=True)
    hyp=OxmlElement('w:hyperlink'); hyp.set(qn('r:id'),rid)
    run=OxmlElement('w:r'); pr=OxmlElement('w:rPr')
    fonts=OxmlElement('w:rFonts'); fonts.set(qn('w:eastAsia'),'맑은 고딕'); fonts.set(qn('w:ascii'),'Malgun Gothic'); fonts.set(qn('w:hAnsi'),'Malgun Gothic'); pr.append(fonts)
    color=OxmlElement('w:color'); color.set(qn('w:val'),'235B80'); pr.append(color)
    sz=OxmlElement('w:sz'); sz.set(qn('w:val'),'22'); pr.append(sz)
    u=OxmlElement('w:u'); u.set(qn('w:val'),'single'); pr.append(u)
    run.append(pr); tx=OxmlElement('w:t'); tx.text=label; run.append(tx); hyp.append(run); par._p.append(hyp)
def slink(needle,label=None):
    e=source(needle); par=doc.add_paragraph(); link(par,label or e['name'],e['url']); return par

def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for col,width in zip(t.columns,widths): col.width=Inches(width)
    tPr=t._tbl.tblPr; borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+edge); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'5'); el.set(qn('w:color'),'D9D9D9'); borders.append(el)
    tPr.append(borders)
    for cells,vals,idx in [(t.rows[0].cells,headers,0)]+[(t.add_row().cells,row,i+1) for i,row in enumerate(rows)]:
        trPr=cells[0]._tc.getparent().get_or_add_trPr()
        if idx==0:
            repeat=OxmlElement('w:tblHeader'); repeat.set(qn('w:val'),'true'); trPr.append(repeat)
        cant=OxmlElement('w:cantSplit'); trPr.append(cant)
        for j,(cell,val,width) in enumerate(zip(cells,vals,widths)):
            cell.width=Inches(width); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=cell._tc.get_or_add_tcPr(); margins=OxmlElement('w:tcMar')
            for side,n in [('top',95),('bottom',95),('left',105),('right',105)]:
                e=OxmlElement('w:'+side); e.set(qn('w:w'),str(n)); e.set(qn('w:type'),'dxa'); margins.append(e)
            pr.append(margins); shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'293D50' if idx==0 else ('F2F5F6' if idx%2==0 else 'FFFFFF')); pr.append(shade)
            par=cell.paragraphs[0]; par.paragraph_format.space_after=Pt(0); par.paragraph_format.space_before=Pt(0); par.paragraph_format.line_spacing=Pt(14.5)
            if isinstance(val,tuple):
                label,url=val; link(par,label,url)
            else:
                run=par.add_run(str(val)); run.font.size=Pt(11); run.bold=idx==0
                if idx==0: run.font.color.rgb=RGBColor(255,255,255)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)
    return t

h('Project M 무료 VFX 조사와 적용 계획',0)
p('2026년 10월 29일  |  Unity 6000.3.21f1  |  URP 17.3.0',bold=True)
p('무료 완성형 툰 VFX를 주력으로 도입하고, 참격과 짧은 잔상만 별도 보완하는 방향을 권장합니다. 우선 비교할 조합은 Cartoon FX Remaster Free와 Free Slash VFX이며, 마력탄 꼬리는 Trails VFX의 URP판을 보조로 검토합니다. Kenney CC0 텍스처는 먼지와 작은 입자를 통일하는 기초 재료로 사용합니다.')
p('이 문서는 스킬 담당자가 최종 프리팹을 받을 수 있도록 후보의 역할, 무료 범위, 도입 비용, 기존 코드와의 접점을 정리합니다. 현재 단계는 배포 페이지와 프로젝트 파일에 근거한 선정 계획입니다. 공식 호환 표기와 우리 프로젝트의 임포트 및 실행 검수 통과는 구분합니다.')
h('권장 결정',2)
bullet('주력은 하나의 툰 패키지로 제한합니다. 폭발과 피격의 색, 명도, 연기 밀도가 효과마다 바뀌지 않게 맞춥니다.')
bullet('참격은 Free Slash VFX와 기존 CombatVFX_v1을 실제 게임 카메라에서 비교합니다. 작은 피격과 폭발은 Simple FX 및 Free Stylized Impacts를 대조군으로 둡니다.')
bullet('전량 직접 제작은 후순위입니다. 필요한 경계 링과 예고 표시만 간단한 메시와 ParticleSystem으로 보완합니다.')
h('아직 결정해야 하는 것',2)
p('마력탄의 기획은 목표 지점 즉시 범위 공격이고 현재 코드는 날아가는 투사체입니다. 이 규칙을 먼저 확정해야 효과의 이동과 착탄 시점을 고정할 수 있습니다. 나머지는 후보 2종을 같은 카메라와 크기로 비교한 뒤 최종 프리팹을 확정하면 됩니다.')
p('최종 선정 조건은 무료 상업 사용 근거, Unity 6와 URP 적용 가능성, 정면 횡스크롤 가독성, 필요한 보정량입니다. 평점과 패키지의 효과 수만으로 결정하지 않습니다.')

page('1 게임 스타일에 맞는 효과')
p('우리 게임은 약 3등신의 로우폴리 3D 주인공과 각진 스켈레톤이 저채도 도시와 옥상에서 싸우는 2.5D 액션입니다. 큰 면과 단순한 실루엣이 중심이므로 VFX도 짧고 선명한 형태를 유지해야 합니다. 픽셀 HUD가 있다는 이유로 월드의 주력 효과를 픽셀아트로 바꾸지는 않습니다.')
img=ROOT/'ArtSources/Stage01Kit/Expansion20261029/Revisions/20261029-01/난간캐릭터검수.png'
doc.add_picture(str(img),width=Inches(6.95))
doc.paragraphs[-1].paragraph_format.line_spacing=1.0
p('현재 주인공과 환경의 비율 및 저채도 팔레트를 확인한 보존 렌더입니다. 효과의 명도와 면적은 이 크기의 캐릭터를 기준으로 비교합니다.')
table(['역할','권장 형태와 색','줄일 표현'],[
 ['물리 공격','아이보리와 옅은 금빛의 호와 짧은 선','큰 오라와 긴 잔광'],
 ['마법 공격','청록 코어와 보라 보조광 및 분명한 꼬리','과한 네온과 다층 광원'],
 ['폭발','짧은 밝은 중심과 각진 파편 및 소량의 연기','실사 볼륨과 화면을 덮는 연무'],
 ['범위와 예고','바닥 경계와 좌우 끝의 짧은 수직 표시','눕힌 원형 문양만으로 범위 전달']
],[1.05,3.6,2.3])
p('MainScene 카메라는 정면 직교 투영이며 orthographic size는 8입니다. 바닥의 XZ 원형 마법진은 옆에서 얇게 보입니다. 폭발은 카메라를 향하는 면으로, 검기는 XY 진행 방향으로 정렬하고, 장판은 경계 양 끝을 세워 폭과 높이를 읽히게 합니다.')

page('2 우선 비교할 무료 패키지')
p('아래 후보는 역할이 겹치더라도 함께 전부 사용하는 목록이 아닙니다. 주력 한 개와 보조 한 개를 먼저 비교하고, 빠진 표현만 추가합니다. 링크는 공식 취득 페이지입니다.')
primary=[
 ('Cartoon FX Remaster','주력 폭발과 피격\n무료 EULA\nUnity 6 URP','굵은 별과 링 및 파편이 적합합니다. 코믹 글자와 과한 Bloom 및 내장 흔들림은 줄입니다. 6.2 API 수정 릴리스가 있습니다.'),
 ('Free Slash VFX','검기와 연속 참격\n무료 EULA\nUnity 6 URP','6.0.34 URP 표기입니다. 초승달 형태와 색을 통일하고 왜곡은 끕니다. 실제 타격 간격에 수명을 맞춥니다.'),
 ('Trails VFX','마력탄과 대시 꼬리\n무료 EULA\nUnity 6 URP','6.0.23 URP 표기이며 원소 꼬리10종입니다. 하나만 선별하고 셰이더와 소리 중복을 줄입니다.'),
 ('Simple FX','먼지와 잔해 대조\n무료 EULA\nURP 가이드','각진 폭발과 돌이 도시 아트에 가깝습니다. Unity 6 직접 표기는 없으며 공식 URP 전환 안내에 따라 확인합니다.'),
 ('Free Stylized Impacts','작은 적중 대조\n무료 EULA\nUnity 6.3 URP','6.3.18 호환 표기와 효과3종입니다. 현재 버전과 가깝지만 실제 비주얼과 성능 검수는 별도로 필요합니다.')
]
rows=[]
for needle,role,note in primary:
    e=source(needle); rows.append([(e['name'],e['url']),role,note])
table(['공식 패키지','프로젝트 역할','선정과 가공 기준'],rows,[2.05,1.5,3.4])
h('도입 순서',2)
p('Cartoon FX Remaster Free와 Simple FX에서 폭발 하나씩, Free Slash VFX와 기존 샘플에서 검기 하나씩을 골라 같은 장면에서 비교합니다. 캐릭터 얼굴과 적의 공격 예고를 가리지 않는 후보를 남기고 그 뒤 마력탄 꼬리와 장판을 맞춥니다.')
p('무료판을 받았다고 전체 유료판의 효과가 제공되는 것은 아닙니다. Standard Unity Asset Store EULA와 제품별 라이선스 표기를 취득 시 보존하고, 에셋 원본을 공개 저장소에 재배포하지 않습니다.')

page('3 보조 재료와 별도 대안')
secondary=[
 ('Particle Pack','Kenney CC0 재료','단순 마스크를 이용해 먼지와 스파크를 통일합니다. 완성 프리팹 대신 방출과 수명을 직접 구성합니다.'),
 ('Smoke Particles','대시와 착지 및 잔연기','낮은 밀도의 먼지에 적합합니다. 큰 연기 덩어리는 줄입니다.'),
 ('Free Animated Explosion','CraftPix 비픽셀 폭발','벡터 기반 큰 형태가 비교 후보입니다. PNG 시트 재생과 알파 및 색 보정이 필요합니다.'),
 ('Hit Animation','OpenGameArt 작은 타격','16프레임 CC0 타격입니다. 텍스처 크기를 줄이고 마지막 타격 또는 작은 적중에 씁니다.'),
 ('Effekseer 공식','편집과 베이크 대안','현 CC0 샘플 구역만 사용합니다. 필요한 효과를 고르고 측면 PNG 시트로 베이크하는 방식을 우선 검토합니다.'),
 ('VFX Mesh Lab','직접 제작 도구','MIT 도구로 초승달과 링 메시를 만듭니다. 기존 메시가 충분하면 추가 도입하지 않아도 됩니다.'),
 ('URP Flipbook VFX Bundle','예비 플립북 비교군','Store판은 무료 EULA와 Unity 6 URP 표기입니다. Lite에 색상 선택과 RGBA 마스크 및 masked emission이 빠져 팔레트 통일에 불리합니다.')
]
rows=[]
for needle,role,note in secondary:
    e=source(needle); rows.append([(e['name'],e['url']),role,note])
table(['공식 배포물','용도','도입 시 할 일'],rows,[2.05,1.5,3.4])
p('Kenney의 단순 마스크와 Effekseer의 편집 가능한 효과는 서로 다른 종류의 대안입니다. 전자는 Unity에서 직접 구성하는 재료이고, 후자는 외부 편집 후 베이크하거나 런타임을 별도로 연결해야 합니다. 이미지 팩을 URP 호환 완성 VFX로 표시하지 않습니다.')

page('4 배포처별 조사 결론')
p('Unity Asset Store 밖에서도 제작자 페이지와 라이선스를 확인했습니다. 같은 팩의 itch.io판과 스토어판은 중복 제품군으로 보며, 무료라는 표시만 있고 상업 조건이 없는 배포판은 조건부로 분류합니다.')
table(['배포처','이번 계획에서의 활용'],[
 ['Unity Asset Store','Unity 프리팹과 공식 파이프라인 표기를 확인하기 쉬운 주력 취득 경로입니다. 최신 Unity 버전은 실제 임포트로 검증합니다.'],
 ['Kenney와 OpenGameArt','CC0 텍스처와 작은 타격 재료의 우선 출처입니다. 모음 제목보다 개별 저작자와 라이선스를 확인합니다.'],
 ['itch.io와 제작자 사이트','무료 툰 팩과 도구의 선택 폭이 큽니다. 0달러 선택 범위와 유료 보너스 및 상업 조건을 구분합니다.'],
 ['CraftPix와 GameDev Market','상업 사용 가능한 무료 PNG 비교군입니다. 원본 재배포 조건과 시트 프레임 구성을 확인합니다.'],
 ['Effekseer와 GitHub','편집 가능한 샘플과 제작 도구를 확보할 수 있습니다. 소프트웨어 라이선스와 효과 콘텐츠의 라이선스는 별개입니다.'],
 ['Gumroad와 GDevelop','무료 텍스처와 픽셀 효과를 추가 조사했습니다. 현재 주력보다 보조 재료와 참고에 가깝습니다.'],
 ['JangaFX','CC0 VDB가 있지만 실사 볼륨이며 베이크 비용이 큽니다. 현재 게임의 주력에서는 제외합니다.'],
 ['Fab와 Godot 배포처','Niagara와 Godot 전용 효과를 Unity에 바로 쓸 수 없습니다. 포팅 또는 베이크 비용 때문에 후순위입니다.'],
 ['BOOTH와 ArtStation','법인 조건 또는 기본 라이선스 한도 등이 있는 제품을 무조건 상업 무료로 분류하지 않았습니다.'],
 ['GameArt2D와 기타 모델 사이트','목표 스킬에 맞는 독립 무료 VFX를 확정하지 못한 목록은 후보 수를 늘리기 위해 넣지 않았습니다.']
],[1.75,5.2])
p('전투 주력에서 제외한 대표 유형은 실사 VDB, 화면 합성용 영상, 픽셀 전용 저해상도 효과, 다른 엔진 전용 그래프, 개인 용도로만 무료인 제품입니다. 각각 무료 여부와 게임 스타일 및 이식 비용을 따로 평가했습니다.')

page('5 직접 제작한다면')
p('직접 만든다면 생성형 3D로 효과마다 모델을 받기보다 ParticleSystem과 간단한 메시를 조합하는 방식이 적합합니다. 현재 프로젝트에 이미 초승달, 링, 파편과 코어가 있고 외부 텍스처가 없는 URP 샘플도 있습니다. 새 제작은 이 자산을 다듬는 범위에서 시작합니다.')
table(['공용 블록','구성','재사용 스킬'],[
 ['참격','초승달 메시와 짧은 선 및 알파 소멸','검기와 회전 및 연속 참격'],
 ['작은 타격','별 모양 섬광과 소수 스파크','검기 적중과 마력탄 착탄'],
 ['마력탄','작은 코어와 짧은 TrailRenderer','마력탄 비행'],
 ['범위','얇은 링과 양 끝 수직 표시','마력장과 폭발 예고'],
 ['폭발','짧은 중심광과 각진 파편 및 소량 연기','일반 폭발과 연쇄 폭발']
],[1.15,3.5,2.3])
h('도구 선택',2)
p('ParticleSystem은 현재 프로젝트에 있는 기본 도구로 제어합니다. URP Particles Unlit을 기본 재질로 두고, 알파와 Additive를 구분합니다. Shader Graph는 공통 dissolve나 텍스처 이동이 필요할 때 한 개부터 추가합니다. VFX Graph는 현재 설치돼 있지 않으며 이 규모의 짧은 효과를 위해 우선 도입할 이유가 작습니다.')
p('Blender는 저면수 초승달과 파편을 편집하거나 툰 플립북을 렌더할 때 사용합니다. Effekseer는 CC0 샘플을 수정해 PNG로 출력하는 별도 경로입니다. 런타임 플러그인을 도입하면 Unity 버전과 빌드 검증이 추가됩니다.')
p('AI 이미지 생성은 모양과 팔레트의 초안에 한정합니다. 프레임 간 연속성, 투명 알파, 발사와 적중 시점, URP 재질, 회수 로직은 별도 제작이 필요합니다. 유료 생성 크레딧을 무료 에셋 비용에 포함하지 않습니다.')
p('일정 견적은 무료팩 비교와 도입 점검 0.5~1일, 스타일 조정과 스킬 연결 1~2일, 검수 0.5~1일입니다. 기존 샘플 전체를 런타임화하면 초안 2~4인일과 추가 다듬기 시간이 필요합니다. 이는 실작업 측정이 아닌 계획용 추정입니다.')

page('6 현재 스킬과 연결할 시점')
p('여섯 스킬은 실제 판정 코드와 프리팹이 이미 있습니다. 오래된 연구 문서의 로그 전용 스킬 설명을 현재 상태로 재사용하지 않습니다. 모든 스킬 SO의 CastDuration은 0.4초이며 효과의 총 재생 시간과 같은 값이 아닙니다.')
table(['스킬','현재 코드와 프리팹','VFX 연결 기준'],[
 ['검기','속도15와 사거리5 및 반경0.5의 관통 투사체','SkillProjectile의 visual child를 교체합니다. 이동은 전투 루트만 담당합니다.'],
 ['회전 베기','박스3.5×2×2와 즉시1회 판정','현재 비어 있는 slashVfx에 짧은 호를 연결합니다. 생성 회전은 identity입니다.'],
 ['연속 참격','5타와 간격0.08초 및 마지막0.32초','매 타 작은 참격을 재생합니다. 현재 slashVfx는 비어 있으며 취소 시 남은 타격이 중단됩니다.'],
 ['마력탄','속도12와 사거리8 및 착탄 반경1.5','현재 투사체형과 기획상 즉시형을 먼저 결정합니다. 폭발 VFX 슬롯은 비어 있습니다.'],
 ['마력장','박스4×1.5×2와 5틱 및 간격0.5초','첫 틱은 즉시이며 마지막은 약2초입니다. 시각 경계와 실제 활성 시간을 맞춥니다.'],
 ['연쇄 폭발','5회와 간격0.12초 및 마지막0.48초','예고와 Detonate를 분리합니다. 예고 슬롯은 비어 있고 현재 폭발은 구체 임시 표현입니다.']
],[1.05,2.9,3.0])
p('마력탄 프리팹의 explodeAtEnd는 현재 false입니다. 사거리 끝에서도 폭발한다는 스크립트의 설정 권고와 실제 직렬화 값은 다릅니다. 최종 규칙을 정한 뒤 시각 효과를 맞춥니다.')
p('연쇄 폭발의 마지막 타격 0.48초는 시전자 CastDuration 0.4초보다 늦습니다. 이미 생성된 폭발이 지연 실행되는 구조이므로 시전자 애니메이션이 끝날 때 모든 폭발 연출을 강제 종료해서는 안 됩니다.')

page('7 기존 샘플을 그대로 붙일 때 생기는 문제')
p('CombatVFX_v1은 실제 메시와 파티클을 움직이는 독립 시연입니다. 런타임 스킬이 요구하는 위치, 시작과 종료의 계약까지 통합된 패키지는 아닙니다. 기존 샘플은 보존하고 런타임용 variant 또는 wrapper를 만듭니다.')
table(['확인한 구조','필요한 보정'],[
 ['활성화로 자동 재생하지 않음','CombatVfxSample.Play를 명시적으로 호출합니다. SkillExplosion의 SetActive만으로는 시작하지 않습니다.'],
 ['Visual Travel 내부 이동','SkillProjectile에 붙이는 버전에서는 내부 이동을 제거합니다. 두 이동이 겹치면 판정과 시각 위치가 달라집니다.'],
 ['샘플 내부의 고정 높이','검기1.05m와 마력탄1.1m 등의 오프셋을 정규화합니다. 스킬 생성 높이에 다시 더하지 않습니다.'],
 ['샘플 폭발 총길이1.55초','현재 폭발 linger0.3초 아래 넣으면 잘립니다. 잔광을 별도 회수하거나 실제 종료까지 수명을 유지합니다.'],
 ['완료 이벤트만 제공','Completed가 발생해도 자동 Destroy 또는 Return은 하지 않습니다. 회수 주체를 하나로 정합니다.'],
 ['프리팹 루트 컴포넌트 차이','typed Spawn에 필요한 SkillProjectile 또는 SkillAutoRelease를 런타임 루트에 배치합니다.'],
 ['장판과 폭발의 루트 배율','장판4×0.1×2와 폭발1.5의 기존 배율 아래 새 효과를 그대로 넣지 않습니다. 단위 루트와 시각 배율을 분리합니다.']
],[2.15,4.8])
p('샘플은 대시0.60초, 검기0.95초, 마력탄1.05초, 폭발1.55초입니다. 이는 시연 길이이며 최종 스킬 지속시간이나 성능 보장이 아닙니다. 입자 상한과 렌더러 수 역시 실제 드로콜 또는 GPU 시간 측정과 다릅니다.')
p('SkillSpawner는 현재 UsePool=false이므로 Instantiate와 Destroy를 사용합니다. 최신 PoolManager는 범용 등록과 비활성 대여를 지원하지만, 먼저 프리팹 등록과 초기화 순서를 처리한 뒤 켜야 합니다.')

page('8 제작과 통합 실행 순서')
h('첫 단계는 같은 화면에서 비교',2)
p('주력 폭발 두 개와 검기 두 개를 같은 카메라, 캐릭터, 배경, 크기로 배치합니다. 패키지 데모의 Bloom과 조명에 의존하지 않고 얼굴 가림, 적 예고와의 겹침, 실루엣을 확인합니다. 이 결과로 주력팩과 보조팩을 확정합니다.')
h('다음 단계는 공통 규격 정리',2)
bullet('루트 scale은 1로 두고 피벗은 타격점 또는 바닥 기준점에 맞춥니다. 아트 크기와 오프셋은 Visual child에서 조정합니다.')
bullet('전방은 로컬 +Z로 통일하고 XY 방향에 LookRotation을 적용합니다. 폭발은 카메라 면을, 검기는 진행 방향을 유지합니다.')
bullet('전투가 위치와 피해 및 충돌을 담당합니다. VFX는 시전과 적중 및 종료 이벤트를 표시합니다.')
bullet('Configure와 Play 및 Stop과 Reset을 구분합니다. 방출을 멈춘 뒤 잔광과 연기가 끝나면 한 번만 회수합니다.')
h('그 뒤 스킬별 변형을 연결',2)
p('최소 전달 프리셋은 대시, 검기, 회전, 연속참격, 마력탄 비행, 마력탄 적중, 마력장, 연쇄 예고, 폭발입니다. 여러 프리셋은 같은 메시와 텍스처를 공유할 수 있습니다. 프리팹마다 전방, 피벗, 반경, 총길이, 잔광길이, 색, 재질 의존성을 한 줄로 기록합니다.')
h('마지막 단계는 재대여와 취소 처리',2)
p('풀에서는 비활성 상태로 대여해 위치와 색을 설정하고, 파티클과 트레일을 지운 뒤 활성화와 Play를 실행합니다. 사망, 피격 취소, 씬 해제 또는 반납 때 예약 콜백과 비동기 작업을 해제합니다. 이전 대여의 늦은 종료 콜백이 새 효과를 회수하지 않도록 세대 토큰을 확인합니다.')
p('스킬 담당자에게 전달할 확정 내용은 패키지 이름만이 아닙니다. 실제 사용할 프리팹과 라이선스 근거, 이벤트 시점, 시작과 종료 API, 좌우 반전, 스케일 및 회수 규칙이 함께 있어야 합니다.')

page('9 무료 범위와 승인 검수')
h('취득과 라이선스',2)
p('Standard Asset Store EULA, CC0, MIT, 제작자 자체 허가는 서로 다릅니다. CC0 텍스처는 표시 의무가 없지만 MIT 도구의 허가문 보존과 무료팩 원본 재배포 금지 조건은 별도로 적용될 수 있습니다. 무료판의 제품 페이지와 받은 버전 및 LICENSE를 함께 보존합니다.')
p('동일 제작자의 itch.io 무료판과 Unity Store판의 조건을 자동으로 합치지 않습니다. 본문에 상업 조건이 없는 판은 정식 스토어판 또는 패키지 라이선스를 확인한 뒤 확정합니다. 팀원이 취득할 경로와 원본의 공개 저장소 포함 여부도 정합니다.')
h('인게임 통과 조건',2)
bullet('실제 정면 카메라에서 좌우, 대각, 공중, 계단, 적 앞뒤와 벽 가림을 확인합니다. 범위선은 실제 판정 폭과 높이에 맞춥니다.')
bullet('일반 공격은 짧게 유지하고 연속 중간 타격은 약하게 합니다. DOT 틱마다 큰 폭발이나 화면 흔들림을 반복하지 않습니다.')
bullet('타격음과 흰 번쩍임 및 피해 숫자를 실제 피해 프레임에 맞춥니다. 기존 UI 타격 스파크와 새 월드 효과의 중복을 줄입니다.')
bullet('PC와 Mobile URP 프로필에서 셰이더 누락, 깊이와 불투명 텍스처 의존, 알파 가장자리, 정렬과 Bloom 의존을 점검합니다.')
bullet('폭발1개와5개 및10개, 겹친 장판, 카메라 이동, 일시정지와 재개, 취소와 재대여를 시험합니다.')
bullet('Development Build에서 CPU와 GPU 및 GC와 투명 오버드로를 기록합니다. 입자 수가 적거나 GPU 방식이라는 이유만으로 빠르다고 판단하지 않습니다.')
p('확정은 임포트와 이 검수를 거친 뒤 수행합니다. 이번 조사에서 공식적으로 표기된 호환성은 후보를 줄이는 근거이며, Unity 6000.3.21f1과 URP17.3.0에서 실행 검수를 통과했다는 의미는 아닙니다.')

page('10 주요 출처와 팀 전달안')
p('제품명 링크는 각 패키지의 공식 취득 페이지입니다. 아래 문서는 제작 방식과 무료 이용 조건을 확인할 때 참고할 1차 출처입니다.')
for label,url in [
 ('Unity Asset Store 이용약관','https://unity.com/legal/as-terms'),
 ('Unity Particle System 선택 기준','https://docs.unity3d.com/6000.3/Documentation/Manual/ChoosingYourParticleSystem.html'),
 ('URP Particles Unlit 재질','https://docs.unity3d.com/6000.3/Documentation/Manual/urp/particles-unlit-shader.html'),
 ('Texture Sheet Animation','https://docs.unity3d.com/6000.3/Documentation/Manual/PartSysTexSheetAnimModule.html'),
 ('Trail Renderer','https://docs.unity3d.com/6000.3/Documentation/Manual/class-TrailRenderer.html'),
 ('CraftPix 무료 제품 라이선스','https://craftpix.net/file-licenses/'),
 ('Effekseer CC0 기여 샘플 구역','https://effekseer.github.io/en/contribute.html'),
 ('Kenney Particle Pack 제작자 배포','https://opengameart.org/content/particle-pack-80-sprites')]:
    par=doc.add_paragraph(); link(par,label,url)
h('프로젝트 근거',2)
p('현재 구현은 MainScene 카메라, Player와 Player Chibi 프리팹, 여섯 PlayerSkillData, SkillProjectile과 SkillField 및 SkillExplosion, SkillSpawner와 PoolManager를 기준으로 확인했습니다. 기존 샘플은 CombatVFX v1의 README와 BuildReport 및 샘플 코드를 대조했습니다. 실제 스킬의 타이밍은 오래된 연구 문서보다 현재 직렬화 값과 코드를 우선합니다.')
h('스킬 담당자에게 전달할 내용',2)
p('“VFX는 무료 툰 팩을 주력으로 하고 참격과 꼬리만 보완하는 방향으로 진행합니다. Cartoon FX Remaster Free와 Free Slash VFX를 우선 비교하고 Trails URP는 마력탄 꼬리 후보로 검토합니다. 같은 게임 카메라에서 색과 크기 및 수명을 맞춘 뒤 최종 프리팹과 Play 및 Stop과 회수 규칙을 전달하겠습니다. 마력탄의 즉시형 또는 투사체형 규칙을 먼저 맞추면 해당 효과도 고정할 수 있습니다.”')
p('상세 무료 후보와 제외 근거는 함께 제공하는 전체 조사 TXT를 참조합니다. 이 보고서는 우선 선택할 조합과 실제 스킬에 붙이기 위한 작업을 중심으로 정리합니다.')

OUT.parent.mkdir(parents=True,exist_ok=True); doc.save(OUT)
print(str(OUT))

if '--render' in sys.argv:
    QA.mkdir(parents=True,exist_ok=True)
    os.environ['PATH']=str(RUNTIME/'native/poppler/Library/bin')+os.pathsep+os.environ.get('PATH','')
    spec=importlib.util.spec_from_file_location('packaged_render_docx',SKILL/'render_docx.py')
    renderer=importlib.util.module_from_spec(spec); spec.loader.exec_module(renderer)
    def word_convert(doc_path,user_profile,convert_tmp_dir,stem,verbose):
        pdf_path=str(Path(convert_tmp_dir)/(stem+'.pdf'))
        def psq(v): return "'"+str(v).replace("'","''")+"'"
        ps="$ErrorActionPreference='Stop'; $w=$null; $d=$null; try { $w=New-Object -ComObject Word.Application; $w.Visible=$false; $w.DisplayAlerts=0; $d=$w.Documents.Open("+psq(doc_path)+",$false,$true); $d.Repaginate(); $d.Fields.Update() | Out-Null; $d.ExportAsFixedFormat("+psq(pdf_path)+",17); } finally { if($d){$d.Close(0)}; if($w){$w.Quit(); [Runtime.InteropServices.Marshal]::ReleaseComObject($w) | Out-Null} }"
        proc=subprocess.run(['powershell','-NoProfile','-NonInteractive','-Command',ps],capture_output=True,text=True,timeout=110)
        log='Native Word COM PDF export\n'+proc.stdout+'\n'+proc.stderr
        (QA/'word-export.log').write_text(log,encoding='utf-8')
        return (pdf_path if Path(pdf_path).exists() else '',log)
    renderer.convert_to_pdf=word_convert
    paths=renderer.rasterize(str(OUT),str(QA),120,verbose=True,emit_pdf=True)
    manifest={'docx':str(OUT),'pages':len(paths),'page_images':list(paths),'renderer':'Packaged render_docx.rasterize; only convert_to_pdf overridden with native Word COM','visually_reviewed':False}
    (QA/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'pages':len(paths),'qa':str(QA)},ensure_ascii=False))
